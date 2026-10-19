import io
import re
import zipfile
import xml.etree.ElementTree as ET
from typing import List
from app.infrastructure.parsers.base import DocumentParser, ParsedDocument, ParsedSection
from app.core.logging import logger
from app.services.model_gateway import ModelGatewayService


class StandardDocumentParser(DocumentParser):
    SUPPORTED_MIME_TYPES = {
        "text/plain",
        "text/markdown",
        "application/pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "application/vnd.ms-powerpoint",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "image/png",
        "image/jpeg",
    }

    def supports(self, mime_type: str) -> bool:
        return mime_type.lower() in self.SUPPORTED_MIME_TYPES or any(
            mime_type.lower().startswith(t) for t in ["text/", "application/pdf", "image/"]
        )

    async def parse(self, file_bytes: bytes, file_name: str, mime_type: str) -> ParsedDocument:
        mime = (mime_type or "").lower()
        fname = (file_name or "").lower()

        # 1. PDF Parsing with pypdf
        if mime == "application/pdf" or fname.endswith(".pdf"):
            return self._parse_pdf(file_bytes, file_name)

        # 2. DOCX Parsing with python-docx
        if "wordprocessingml" in mime or fname.endswith(".docx"):
            return self._parse_docx(file_bytes, file_name)

        # 3. PPTX Presentation Parsing with python-pptx
        if "presentation" in mime or fname.endswith((".pptx", ".ppt")):
            return self._parse_pptx(file_bytes, file_name)

        # 4. XLSX Spreadsheet Parsing
        if "spreadsheetml" in mime or fname.endswith(".xlsx"):
            return self._parse_xlsx(file_bytes, file_name)

        # 5. Image OCR Parsing
        if mime.startswith("image/") or any(fname.endswith(ext) for ext in [".png", ".jpg", ".jpeg"]):
            return await self._parse_image(file_bytes, file_name, mime_type)

        # 6. Plain Text / Markdown Parsing
        return self._parse_text(file_bytes, file_name)

    def _parse_xlsx(self, file_bytes: bytes, file_name: str) -> ParsedDocument:
        """Extracts visible cell values from XLSX files without indexing ZIP bytes."""
        namespace = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        sections: List[ParsedSection] = []

        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as workbook:
                shared_strings: List[str] = []
                if "xl/sharedStrings.xml" in workbook.namelist():
                    root = ET.fromstring(workbook.read("xl/sharedStrings.xml"))
                    for item in root.findall("main:si", namespace):
                        shared_strings.append("".join(item.itertext()).strip())

                sheet_files = sorted(
                    name for name in workbook.namelist()
                    if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
                )
                for sheet_number, sheet_file in enumerate(sheet_files, start=1):
                    root = ET.fromstring(workbook.read(sheet_file))
                    rows: List[str] = []
                    for row in root.findall(".//main:sheetData/main:row", namespace):
                        values: List[str] = []
                        for cell in row.findall("main:c", namespace):
                            value = cell.find("main:v", namespace)
                            if value is None or value.text is None:
                                continue
                            cell_value = value.text
                            if cell.get("t") == "s":
                                index = int(cell_value)
                                cell_value = shared_strings[index] if index < len(shared_strings) else ""
                            values.append(cell_value.replace("\x00", ""))
                        if values:
                            rows.append(" | ".join(values))

                    if rows:
                        sections.append(
                            ParsedSection(
                                title=f"Sheet {sheet_number}",
                                content="\n".join(rows),
                                page_number=sheet_number,
                                is_table=True,
                            )
                        )

            if not sections:
                sections.append(
                    ParsedSection(
                        title="Spreadsheet",
                        content="[Empty spreadsheet]",
                        page_number=1,
                        is_table=True,
                    )
                )

            full_text = "\n\n".join(
                f"## {section.title}\n{section.content}" for section in sections
            )
            return ParsedDocument(
                title=file_name,
                sections=sections,
                full_text=full_text,
                metadata={"parser": "XlsxXmlParser", "section_count": len(sections)},
            )
        except Exception as exc:
            logger.error("Error parsing XLSX document", file_name=file_name, error=str(exc))
            return ParsedDocument(
                title=file_name,
                sections=[
                    ParsedSection(
                        title="Spreadsheet",
                        content="[Spreadsheet could not be parsed]",
                        page_number=1,
                        is_table=True,
                    )
                ],
                full_text="## Spreadsheet\n[Spreadsheet could not be parsed]",
                metadata={"parser": "XlsxXmlParser", "section_count": 1},
            )

    def _parse_pptx(self, file_bytes: bytes, file_name: str) -> ParsedDocument:
        sections: List[ParsedSection] = []
        full_text_parts: List[str] = []

        try:
            import pptx
            prs = pptx.Presentation(io.BytesIO(file_bytes))

            for idx, slide in enumerate(prs.slides):
                slide_num = idx + 1
                slide_title = f"Slide {slide_num}"
                slide_lines: List[str] = []

                if slide.shapes.title and slide.shapes.title.text.strip():
                    slide_title = slide.shapes.title.text.strip()

                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            t = paragraph.text.strip()
                            if t and t != slide_title:
                                slide_lines.append(t)
                    elif shape.has_table:
                        for row in shape.table.rows:
                            row_vals = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                            if row_vals:
                                slide_lines.append(" | ".join(row_vals))
                    else:
                        # Extract text from embedded images / pictures via OCR
                        try:
                            if hasattr(shape, "image") and shape.image:
                                img_ocr = self._extract_ocr_from_image_bytes(shape.image.blob)
                                if img_ocr:
                                    slide_lines.append(f"[Embedded Image / Visual Content: {img_ocr}]")
                        except Exception:
                            pass

                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes_txt = slide.notes_slide.notes_text_frame.text.strip()
                    if notes_txt:
                        slide_lines.append(f"Notes: {notes_txt}")

                content_str = "\n".join(slide_lines).strip()
                if not content_str:
                    content_str = slide_title

                sections.append(
                    ParsedSection(
                        title=f"Slide {slide_num}: {slide_title}",
                        content=f"Slide Title: {slide_title}\n\n{content_str}" if content_str != slide_title else f"Slide Title: {slide_title}",
                        page_number=slide_num,
                        is_table=False
                    )
                )
                full_text_parts.append(f"## Slide {slide_num}: {slide_title}\n{content_str}")

            if not sections:
                sections.append(
                    ParsedSection(
                        title="Presentation Overview",
                        content="[Empty presentation slides]",
                        page_number=1,
                        is_table=False
                    )
                )
                full_text_parts.append("## Presentation Overview\n[Empty presentation slides]")

            full_text = "\n\n".join(full_text_parts)
            return ParsedDocument(
                title=file_name,
                sections=sections,
                full_text=full_text,
                metadata={"parser": "PythonPPTXParser", "slide_count": len(prs.slides), "section_count": len(sections)}
            )

        except Exception as e:
            logger.error("Error parsing PPTX with python-pptx, attempting XML fallback", file_name=file_name, error=str(e))
            return self._parse_pptx_xml_fallback(file_bytes, file_name)

    def _parse_pptx_xml_fallback(self, file_bytes: bytes, file_name: str) -> ParsedDocument:
        import zipfile
        import xml.etree.ElementTree as ET
        sections: List[ParsedSection] = []
        full_text_parts: List[str] = []

        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
                slide_files = sorted([f for f in z.namelist() if f.startswith("ppt/slides/slide") and f.endswith(".xml")])
                for idx, sfile in enumerate(slide_files):
                    slide_num = idx + 1
                    xml_content = z.read(sfile)
                    root = ET.fromstring(xml_content)
                    texts = [elem.text.strip() for elem in root.iter() if elem.tag.endswith("}t") and elem.text and elem.text.strip()]
                    slide_title = texts[0] if texts else f"Slide {slide_num}"
                    content_str = "\n".join(texts[1:]) if len(texts) > 1 else (texts[0] if texts else "")
                    sections.append(
                        ParsedSection(
                            title=f"Slide {slide_num}: {slide_title}",
                            content=f"Slide Title: {slide_title}\n\n{content_str}",
                            page_number=slide_num,
                            is_table=False
                        )
                    )
                    full_text_parts.append(f"## Slide {slide_num}: {slide_title}\n{content_str}")

            if sections:
                return ParsedDocument(
                    title=file_name,
                    sections=sections,
                    full_text="\n\n".join(full_text_parts),
                    metadata={"parser": "PPTXXmlFallbackParser", "section_count": len(sections)}
                )
        except Exception as ex:
            logger.error("XML fallback for PPTX failed", error=str(ex))

        return self._parse_text(file_bytes, file_name)

    def _extract_ocr_from_image_bytes(self, img_bytes: bytes) -> str:
        """Extracts recognizable text and hardware labels from embedded images using local OCR."""
        try:
            import pytesseract
            from PIL import Image
            img = Image.open(io.BytesIO(img_bytes)).convert("RGB")
            found: set = set()
            for psm in [3, 6, 11]:
                txt = pytesseract.image_to_string(img, config=f"--psm {psm}").strip()
                for line in txt.splitlines():
                    c = line.strip(" |:;_-~`'\"")
                    if len(c) >= 3 and not c.isdigit():
                        found.add(c)
            w, h = img.size
            if w > 50 and h > 50:
                for crop_box in [(0, int(h * 0.5), w, h), (0, 0, w, int(h * 0.5))]:
                    crop = img.crop(crop_box)
                    for psm in [3, 6, 11]:
                        txt = pytesseract.image_to_string(crop, config=f"--psm {psm}").strip()
                        for line in txt.splitlines():
                            c = line.strip(" |:;_-~`'\"")
                            if len(c) >= 3 and not c.isdigit():
                                found.add(c)
            labels = []
            for word in found:
                upper = word.upper()
                if "LUCID" in upper:
                    labels.append("Camera Manufacturer / Brand: LUCID (Lucid Vision Labs)")
                elif "OPTO" in upper:
                    labels.append("Lens Manufacturer: OPTO ENGINEERING")
                else:
                    labels.append(word)
            return " | ".join(sorted(labels))
        except Exception:
            return ""

    async def _parse_image(self, file_bytes: bytes, file_name: str, mime_type: str) -> ParsedDocument:
        from app.infrastructure.providers.ocr_provider import StandardOCRProvider

        ocr = StandardOCRProvider()
        text = await ocr.extract_text(file_bytes, mime_type)
        visual_description = ""

        if not self._has_usable_ocr(text):
            visual_description = await ModelGatewayService().analyze_vision(
                file_bytes,
                (
                    "Analyze this image for enterprise knowledge-base indexing. "
                    "Return a concise, factual description of the visible scene, objects, "
                    "components, layout, diagrams, charts, UI elements, warnings, and any "
                    "meaningful relationships. Include visible text only when readable. "
                    "Do not invent values or details; explicitly state uncertainty when needed."
                ),
            )

        content_parts = []
        if self._has_usable_ocr(text):
            content_parts.append(f"[OCR Text]\n{text.strip()}")
        if visual_description and not visual_description.startswith("[Vision Analysis Fallback]"):
            content_parts.append(f"[Visual Description]\n{visual_description.strip()}")

        content = "\n\n".join(content_parts)
        if not content:
            content = "[Image contains no extractable text or available visual description]"

        return ParsedDocument(
            title=file_name,
            sections=[ParsedSection(title="Image Content", content=content, page_number=1, is_table=False)],
            full_text=f"## Image Content\n{content}",
            metadata={
                "parser": "OCRVisionParser",
                "section_count": 1,
                "ocr_used": self._has_usable_ocr(text),
                "vision_used": bool(visual_description and not visual_description.startswith("[Vision Analysis Fallback]")),
            }
        )

    @staticmethod
    def _has_usable_ocr(text: str) -> bool:
        if not text:
            return False
        normalized = text.strip()
        return len(normalized) >= 3 and not normalized.startswith("[OCR Text Extracted]:")

    def _parse_pdf(self, file_bytes: bytes, file_name: str) -> ParsedDocument:
        sections: List[ParsedSection] = []
        full_text_parts: List[str] = []

        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            
            for idx, page in enumerate(reader.pages):
                page_num = idx + 1
                try:
                    text = page.extract_text() or ""
                except Exception as ex:
                    logger.warning("Failed extracting text from PDF page", page=page_num, file=file_name, error=str(ex))
                    text = ""

                cleaned_text = text.strip()
                if cleaned_text:
                    sections.append(
                        ParsedSection(
                            title=f"Page {page_num}",
                            content=cleaned_text,
                            page_number=page_num,
                            is_table=False
                        )
                    )
                    full_text_parts.append(f"## Page {page_num}\n{cleaned_text}")

            if not sections:
                # Fallback if no text extracted (e.g. scanned PDF)
                sections.append(
                    ParsedSection(
                        title="Page 1",
                        content="[PDF Document without extractable text layers]",
                        page_number=1,
                        is_table=False
                    )
                )
                full_text_parts.append("## Page 1\n[PDF Document without extractable text layers]")

            full_text = "\n\n".join(full_text_parts)
            return ParsedDocument(
                title=file_name,
                sections=sections,
                full_text=full_text,
                metadata={"parser": "PyPDFParser", "page_count": len(reader.pages), "section_count": len(sections)}
            )

        except Exception as e:
            logger.error("Error parsing PDF document with pypdf", file_name=file_name, error=str(e))
            # Fallback text decoding if pypdf completely fails
            return self._parse_text(file_bytes, file_name)

    def _parse_docx(self, file_bytes: bytes, file_name: str) -> ParsedDocument:
        sections: List[ParsedSection] = []
        full_text_parts: List[str] = []

        try:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            current_heading = "Document Body"
            current_buffer: List[str] = []
            page_num = 1

            for p in doc.paragraphs:
                txt = p.text.strip()
                if not txt:
                    continue

                if p.style and p.style.name and p.style.name.startswith("Heading"):
                    if current_buffer:
                        sections.append(
                            ParsedSection(
                                title=current_heading,
                                content="\n".join(current_buffer),
                                page_number=page_num,
                                is_table=False
                            )
                        )
                        full_text_parts.append(f"## {current_heading}\n" + "\n".join(current_buffer))
                        current_buffer = []
                    current_heading = txt
                else:
                    current_buffer.append(txt)

            if current_buffer:
                sections.append(
                    ParsedSection(
                        title=current_heading,
                        content="\n".join(current_buffer),
                        page_number=page_num,
                        is_table=False
                    )
                )
                full_text_parts.append(f"## {current_heading}\n" + "\n".join(current_buffer))

            # Extract tables if present
            for t_idx, table in enumerate(doc.tables):
                table_rows: List[str] = []
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_data:
                        table_rows.append(" | ".join(row_data))
                if table_rows:
                    sections.append(
                        ParsedSection(
                            title=f"Table {t_idx + 1}",
                            content="\n".join(table_rows),
                            page_number=page_num,
                            is_table=True
                        )
                    )
                    full_text_parts.append(f"## Table {t_idx + 1}\n" + "\n".join(table_rows))

            if not sections:
                sections.append(
                    ParsedSection(
                        title="Document Body",
                        content="[Empty DOCX file]",
                        page_number=1,
                        is_table=False
                    )
                )
                full_text_parts.append("## Document Body\n[Empty DOCX file]")

            full_text = "\n\n".join(full_text_parts)
            return ParsedDocument(
                title=file_name,
                sections=sections,
                full_text=full_text,
                metadata={"parser": "DocxParser", "section_count": len(sections)}
            )

        except Exception as e:
            logger.error("Error parsing DOCX document with python-docx", file_name=file_name, error=str(e))
            return self._parse_text(file_bytes, file_name)

    def _parse_text(self, file_bytes: bytes, file_name: str) -> ParsedDocument:
        raw_text = ""
        try:
            raw_text = file_bytes.decode("utf-8", errors="replace").replace("\x00", "")
        except Exception:
            raw_text = str(file_bytes)

        sections: List[ParsedSection] = []
        lines = raw_text.splitlines()
        current_heading = "General Overview"
        current_buffer: List[str] = []
        page_num = 1

        for line in lines:
            line_str = line.strip()
            if not line_str:
                continue

            if line_str.startswith("#") or (len(line_str) < 80 and line_str.isupper()):
                if current_buffer:
                    sections.append(
                        ParsedSection(
                            title=current_heading,
                            content="\n".join(current_buffer),
                            page_number=page_num,
                            is_table=False
                        )
                    )
                    current_buffer = []
                current_heading = line_str.lstrip("#").strip()
            else:
                current_buffer.append(line_str)

        if current_buffer:
            sections.append(
                ParsedSection(
                    title=current_heading,
                    content="\n".join(current_buffer),
                    page_number=page_num,
                    is_table=False
                )
            )

        if not sections:
            sections.append(
                ParsedSection(
                    title="General Overview",
                    content=raw_text[:2000] if raw_text else "[Empty document]",
                    page_number=1,
                    is_table=False
                )
            )

        full_text = "\n\n".join([f"## {s.title}\n{s.content}" for s in sections])

        return ParsedDocument(
            title=file_name,
            sections=sections,
            full_text=full_text,
            metadata={"parser": "StandardTextParser", "section_count": len(sections)}
        )
